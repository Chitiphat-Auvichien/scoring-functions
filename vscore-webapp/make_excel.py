#!/usr/bin/env python3
"""make_excel.py -- one comprehensive workbook of every scored mode.

Scores every molecule in scoring-functions/data that has both a .log and a
.com deck, through the same pipeline the webapp uses, and writes a workbook
with all seven scores per mode, the dominant score highlighted, and the label.

The pipeline is verified against the published *_normal.csv results at
0.00e+00 deviation (tests/test_regression.py); this script additionally
cross-checks its output against data/results/library_scores.csv and writes the
comparison into a Validation sheet, so the workbook carries its own audit.

    python make_excel.py -o "../Vibrational scores (all molecules).xlsx"
"""

from __future__ import annotations

import argparse
import collections
import glob
import os
import re
import sys
from pathlib import Path

import numpy as np
import pandas as pd
from openpyxl.formatting.rule import CellIsRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

sys.path.insert(0, str(Path(__file__).resolve().parent))

from app.core.classifier import Thresholds                       # noqa: E402
from app.core.parsers import (ParseError, parse_connectivity,     # noqa: E402
                              parse_gaussian_log)
from app.core.pipeline import SCORE_KEYS, analyse                # noqa: E402

REF = Path(__file__).resolve().parent.parent / "scoring-functions" / "data"

# ---- palette (readable in Excel's default light theme) ----------------
HDR_FILL = PatternFill("solid", fgColor="1F3B36")
HDR_FONT = Font(color="FFFFFF", bold=True, size=10)
MAX_FILL = PatternFill("solid", fgColor="FFE9A8")     # dominant |score|
MAX_FONT = Font(bold=True, color="6B4B00")
LABEL_FILL = {
    "S":  PatternFill("solid", fgColor="D8ECE7"),
    "B":  PatternFill("solid", fgColor="F6E3D3"),
    "SB": PatternFill("solid", fgColor="E2DFF4"),
}
LABEL_FONT = {
    "S":  Font(bold=True, color="0B5F52"),
    "B":  Font(bold=True, color="8A4B1F"),
    "SB": Font(bold=True, color="4A3F8F"),
}
REF_FILL = PatternFill("solid", fgColor="EFEFEA")
THIN = Side(style="thin", color="D5D5CC")
BORDER = Border(bottom=THIN)


# ======================================================================
# Data collection
# ======================================================================
def point_group(log_path: Path) -> str:
    """Gaussian's 'Full point group' line, normalised (D*H -> D∞h)."""
    with open(log_path, errors="replace") as f:
        for line in f:
            if "Full point group" in line:
                pg = line.split()[3]
                return {"D*H": "D∞h", "C*V": "C∞v"}.get(pg, pg)
    return "?"


def route_line(com_path: Path) -> str:
    with open(com_path, errors="replace") as f:
        for line in f:
            if line.strip().startswith("#"):
                return line.strip()
    return "?"


def level_of_theory(route: str) -> str:
    """The method/basis token out of a Gaussian route line."""
    m = re.search(r"\b((?:mp2|b3lyp|hf|ccsd\(t\)|ccsd|pbe0|m06)\s*/\s*\S+)", route, re.I)
    return m.group(1).replace(" ", "") if m else "?"


def collect():
    rows, refs, bonds, per_mol, failures = [], [], [], [], []
    stems = sorted(
        os.path.basename(p)[:-4] for p in glob.glob(str(REF / "logs" / "*.log")))

    for stem in stems:
        log = REF / "logs" / f"{stem}.log"
        com = next((REF / "gjf" / f"{stem}{e}"
                    for e in (".com", ".gjf", ".inp")
                    if (REF / "gjf" / f"{stem}{e}").exists()), None)
        if com is None:
            failures.append((stem, "no .com/.gjf deck (connectivity is never inferred)"))
            continue
        try:
            g = parse_gaussian_log(log.read_text(errors="replace"))
            bl = parse_connectivity(com.read_text(errors="replace"), len(g["atoms"]))
            p = analyse(g["atoms"], g["coords"], bl, g["modes"], title=stem)
        except ParseError as exc:
            failures.append((stem, str(exc)[:150]))
            continue

        pg, route = point_group(log), route_line(com)
        lot = level_of_theory(route)
        common = {"Molecule": stem, "Point group": pg, "Level of theory": lot}

        for r in p["vibrations"]:
            rows.append({**common, "Mode": r["name"],
                         "Freq / cm-1": r["frequency"],
                         **{k: r["scores"][k] for k in SCORE_KEYS},
                         "Max score": r["dominant"],
                         "Max |value|": abs(r["dominant_value"]),
                         "Label": r["label"], "Meaning": r["label_text"],
                         "Irrep": r["irrep"], "Mu / AMU": r["reduced_mass"],
                         "K / mDyne A-1": r["force_constant"],
                         "Annotation": r["annotation"] or ""})
            for b in r["bonds"]:
                bonds.append({**common, "Mode": r["name"],
                              "Freq / cm-1": r["frequency"], "Label": r["label"],
                              "Bond": b["pair"], "s_AB": b["s_AB"],
                              "rel_db": b["rel_db"]})

        for r in p["references"]:
            refs.append({**common, "Reference mode": r["name"],
                         **{k: r["scores"][k] for k in SCORE_KEYS},
                         "Max score": r["dominant"], "Label": r["label"]})

        c = collections.Counter(r["label"] for r in p["vibrations"])
        d = collections.Counter(r["dominant"] for r in p["vibrations"])
        per_mol.append({
            "Molecule": stem, "Point group": pg, "Level of theory": lot,
            "Atoms": p["n_atoms"], "Bonds": len(p["bonds"]),
            "Modes": p["n_modes"], "Linear": "yes" if p["linear"] else "no",
            "S": c.get("S", 0), "B": c.get("B", 0), "SB": c.get("SB", 0),
            "External": sum(v for k, v in c.items() if k not in ("S", "B", "SB")),
            **{f"max={k}": d.get(k, 0) for k in SCORE_KEYS},
            "Displacement dp": p["precision_dp"],
            "Route": route,
        })

    # Ordered by point group then molecule -- the layout these deliverables
    # have used before, and the autofilter still allows any other view.
    def pg_sort(df, extra=()):
        if df.empty:
            return df
        return df.sort_values(["Point group", "Molecule", *extra],
                              kind="stable").reset_index(drop=True)

    return (pg_sort(pd.DataFrame(rows)), pg_sort(pd.DataFrame(refs)),
            pg_sort(pd.DataFrame(bonds)), pg_sort(pd.DataFrame(per_mol)),
            failures)


def by_point_group(per_mol: pd.DataFrame, modes: pd.DataFrame) -> pd.DataFrame:
    """One row per point group: how many molecules, modes, and labels."""
    out = []
    for pg, grp in per_mol.groupby("Point group"):
        md = modes[modes["Point group"] == pg]
        c = collections.Counter(md["Label"])
        d = collections.Counter(md["Max score"])
        out.append({
            "Point group": pg,
            "Molecules": len(grp),
            "Modes": len(md),
            "S": c.get("S", 0), "B": c.get("B", 0), "SB": c.get("SB", 0),
            **{f"max={k}": d.get(k, 0) for k in SCORE_KEYS},
            "Levels of theory": ", ".join(sorted(grp["Level of theory"].unique())),
            "Molecules listed": ", ".join(sorted(grp["Molecule"])),
        })
    return pd.DataFrame(out).sort_values("Molecules", ascending=False).reset_index(drop=True)


def validate(modes: pd.DataFrame):
    """Compare against the published library_scores.csv, per molecule."""
    path = REF / "results" / "library_scores.csv"
    if not path.exists():
        return pd.DataFrame()
    lib = pd.read_csv(path)
    lib = lib[lib["kind"] == "internal"]
    out = []
    for mol, grp in modes.groupby("Molecule"):
        ref = lib[lib["molecule"] == mol]
        if ref.empty:
            out.append({"Molecule": mol, "In library": "no", "Modes here": len(grp),
                        "Modes in library": 0, "Worst |diff|": None, "Verdict": "not in library"})
            continue
        a = grp.sort_values("Freq / cm-1")
        b = ref.sort_values("freq")
        n = min(len(a), len(b))
        worst = 0.0
        for col, refcol in [("Tx", "Tx"), ("Ty", "Ty"), ("Tz", "Tz"),
                            ("Rx", "Rx"), ("Ry", "Ry"), ("Rz", "Rz"),
                            ("V_S", "V_Stretch")]:
            if refcol in b.columns:
                worst = max(worst, float(np.nanmax(np.abs(
                    a[col].to_numpy()[:n] - b[refcol].to_numpy()[:n]))))
        out.append({"Molecule": mol, "In library": "yes", "Modes here": len(grp),
                    "Modes in library": len(b), "Worst |diff|": round(worst, 6),
                    "Verdict": "match" if worst < 5e-4 and len(a) == len(b)
                               else ("count differs" if len(a) != len(b) else "DIFFERS")})
    return pd.DataFrame(out)


# ======================================================================
# Workbook
# ======================================================================
def style_sheet(ws, df, freeze="A2", widths=None, number_formats=None):
    for c in range(1, len(df.columns) + 1):
        cell = ws.cell(row=1, column=c)
        cell.fill, cell.font = HDR_FILL, HDR_FONT
        cell.alignment = Alignment(horizontal="center", vertical="center",
                                   wrap_text=True)
    ws.freeze_panes = freeze
    ws.auto_filter.ref = ws.dimensions
    ws.row_dimensions[1].height = 30
    for i, col in enumerate(df.columns, 1):
        w = (widths or {}).get(col)
        if w is None:
            w = max(9, min(24, int(df[col].astype(str).str.len().max() or 8) + 2,
                           max(len(str(col)) + 3, 9)))
            w = max(w, len(str(col)) + 3)
        ws.column_dimensions[get_column_letter(i)].width = w
        fmt = (number_formats or {}).get(col)
        if fmt:
            for r in range(2, len(df) + 2):
                ws.cell(row=r, column=i).number_format = fmt


def write(out_path: Path, modes, refs, bonds, per_mol, pgroups, valid, failures):
    thr = Thresholds.calibrated()

    # ---------------- Summary (built as a plain sheet, not a frame) -----
    with pd.ExcelWriter(out_path, engine="openpyxl") as xl:
        summary_rows = [
            ("Vibrational mode scoring — all molecules", ""),
            ("", ""),
            ("Molecules scored", len(per_mol)),
            ("Vibrational modes scored", len(modes)),
            ("Ideal T/R reference modes", len(refs)),
            ("Per-bond s_AB entries", len(bonds)),
            ("Molecules skipped", len(failures)),
            ("", ""),
            ("THE SEVEN SCORES", ""),
            ("s[Tx], s[Ty], s[Tz]", "translation along each principal axis, range [-1, 1]"),
            ("s[Rx], s[Ry], s[Rz]", "rotation about each principal axis, range [-1, 1]"),
            ("s[V_S]", "stretch character, range [0, 1]; 1 = pure bond-length change"),
            ("", ""),
            ("LABELS", ""),
            ("S", f"stretching — s[V_S] >= tau_S = {thr.tau_S:.4f}"),
            ("B", f"bending — s[V_S] <= tau_B = {thr.tau_B:.4f}"),
            ("SB", "mixed stretch/bend — between the two"),
            ("Tx..Rz", "external mode (translation/rotation); a trailing * = mixed with vibration"),
            ("", ""),
            ("HIGHLIGHTING", ""),
            ("Amber cell", "the largest |score| of the seven in that row"),
            ("Why absolute value", "s[V_S] alone cannot be negative, so a raw comparison "
                                   "favours it structurally (64% vs 57% of modes)"),
            ("", ""),
            ("METHOD", ""),
            ("V-score bond weighting", thr.v_weighting),
            ("tau_TR / tau_S / tau_B", f"{thr.tau_TR} / {thr.tau_S:.6f} / {thr.tau_B:.6f}"),
            ("Frame", "molecule and displacements rotated together into principal axes "
                      "of inertia; s[V_S] is rotation-invariant, T and R are not"),
            ("Connectivity", "taken from each molecule's .com deck — never inferred "
                             "from interatomic distances"),
            ("Verified", "reproduces the published *_normal.csv results at 0.00e+00; "
                         "see the Validation sheet"),
            ("", ""),
            ("!! READ THIS", ""),
            ("Level of theory is NOT uniform",
             "the set spans several methods/bases — see the 'Level of theory' column "
             "on By molecule. Do not compare absolute frequencies across molecules "
             "without checking it."),
            ("Real modes never get T/R labels",
             "the 6 ideal reference modes always win the external slots, so every real "
             "vibration is S/B/SB. The Reference modes sheet is an alignment check, "
             "not a result."),
        ]
        pd.DataFrame(summary_rows, columns=["Item", "Value"]).to_excel(
            xl, sheet_name="Summary", index=False)

        modes.to_excel(xl, sheet_name="All modes", index=False)
        per_mol.to_excel(xl, sheet_name="By molecule", index=False)
        pgroups.to_excel(xl, sheet_name="By point group", index=False)
        refs.to_excel(xl, sheet_name="Reference modes", index=False)
        bonds.to_excel(xl, sheet_name="Per-bond s_AB", index=False)
        if not valid.empty:
            valid.to_excel(xl, sheet_name="Validation", index=False)
        if failures:
            pd.DataFrame(failures, columns=["Molecule", "Reason"]).to_excel(
                xl, sheet_name="Skipped", index=False)

        wb = xl.book

        # ---------------- Summary styling ------------------------------
        ws = wb["Summary"]
        ws.column_dimensions["A"].width = 32
        ws.column_dimensions["B"].width = 96
        ws["A1"].font = Font(bold=True, size=14, color="1F3B36")
        for r in range(1, len(summary_rows) + 2):
            ws.cell(row=r, column=2).alignment = Alignment(wrap_text=True, vertical="top")
            a = ws.cell(row=r, column=1)
            if a.value and str(a.value).isupper() and not ws.cell(row=r, column=2).value:
                a.font = Font(bold=True, size=11, color="1F3B36")
            if str(a.value).startswith("!!"):
                a.font = Font(bold=True, size=11, color="9C2B2B")
        ws.sheet_view.showGridLines = False

        # ---------------- All modes ------------------------------------
        ws = wb["All modes"]
        cols = list(modes.columns)
        style_sheet(ws, modes, freeze="E2",
                    widths={"Molecule": 15, "Meaning": 24, "Annotation": 18,
                            "Level of theory": 15, "Max score": 11},
                    number_formats={**{k: "0.0000" for k in SCORE_KEYS},
                                    "Freq / cm-1": "0.00", "Max |value|": "0.0000",
                                    "Mu / AMU": "0.0000", "K / mDyne A-1": "0.0000"})
        i_score = {k: cols.index(k) + 1 for k in SCORE_KEYS}
        i_max, i_label = cols.index("Max score") + 1, cols.index("Label") + 1
        for r in range(2, len(modes) + 2):
            dom = ws.cell(row=r, column=i_max).value
            if dom in i_score:
                c = ws.cell(row=r, column=i_score[dom])
                c.fill, c.font = MAX_FILL, MAX_FONT
            lab = ws.cell(row=r, column=i_label)
            if lab.value in LABEL_FILL:
                lab.fill, lab.font = LABEL_FILL[lab.value], LABEL_FONT[lab.value]
            lab.alignment = Alignment(horizontal="center")
            ws.cell(row=r, column=i_max).alignment = Alignment(horizontal="center")

        # ---------------- By molecule ---------------------------------
        ws = wb["By molecule"]
        style_sheet(ws, per_mol, freeze="D2",
                    widths={"Molecule": 15, "Route": 52, "Level of theory": 16})

        ws = wb["By point group"]
        style_sheet(ws, pgroups, freeze="B2",
                    widths={"Point group": 13, "Levels of theory": 34,
                            "Molecules listed": 80})

        # ---------------- Reference modes -----------------------------
        ws = wb["Reference modes"]
        style_sheet(ws, refs, freeze="E2", widths={"Molecule": 15},
                    number_formats={k: "0.0000" for k in SCORE_KEYS})
        rcols = list(refs.columns)
        ri_max = rcols.index("Max score") + 1
        ri_score = {k: rcols.index(k) + 1 for k in SCORE_KEYS}
        for r in range(2, len(refs) + 2):
            for c in range(1, len(rcols) + 1):
                ws.cell(row=r, column=c).fill = REF_FILL
            dom = ws.cell(row=r, column=ri_max).value
            if dom in ri_score:
                cell = ws.cell(row=r, column=ri_score[dom])
                cell.fill, cell.font = MAX_FILL, MAX_FONT

        # ---------------- Per-bond ------------------------------------
        ws = wb["Per-bond s_AB"]
        style_sheet(ws, bonds, freeze="E2", widths={"Molecule": 15, "Bond": 12},
                    number_formats={"s_AB": "0.0000", "rel_db": "0.00000",
                                    "Freq / cm-1": "0.00"})
        col = get_column_letter(list(bonds.columns).index("s_AB") + 1)
        rng = f"{col}2:{col}{len(bonds) + 1}"
        ws.conditional_formatting.add(rng, CellIsRule(
            operator="greaterThan", formula=["0.05"],
            fill=PatternFill("solid", fgColor="D8ECE7")))
        ws.conditional_formatting.add(rng, CellIsRule(
            operator="lessThan", formula=["-0.05"],
            fill=PatternFill("solid", fgColor="F6E3D3")))

        # ---------------- Validation ----------------------------------
        if not valid.empty:
            ws = wb["Validation"]
            style_sheet(ws, valid, widths={"Molecule": 15, "Verdict": 16},
                        number_formats={"Worst |diff|": "0.000000"})
            iv = list(valid.columns).index("Verdict") + 1
            for r in range(2, len(valid) + 2):
                c = ws.cell(row=r, column=iv)
                if c.value == "match":
                    c.fill, c.font = LABEL_FILL["S"], LABEL_FONT["S"]
                elif c.value and c.value != "match":
                    c.fill, c.font = LABEL_FILL["B"], LABEL_FONT["B"]

        if failures:
            style_sheet(wb["Skipped"], pd.DataFrame(failures,
                        columns=["Molecule", "Reason"]),
                        widths={"Molecule": 16, "Reason": 90})


def main(argv=None):
    ap = argparse.ArgumentParser(description=__doc__,
                                 formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("-o", "--out", default="Vibrational scores (all molecules).xlsx")
    args = ap.parse_args(argv)

    print("scoring every molecule with a .log and a .com deck ...")
    modes, refs, bonds, per_mol, failures = collect()
    print(f"  {len(per_mol)} molecules, {len(modes)} vibrational modes, "
          f"{len(bonds)} per-bond entries, {len(failures)} skipped")

    pgroups = by_point_group(per_mol, modes)
    print(f"  {len(pgroups)} point groups: "
          + ", ".join(f"{r['Point group']}({r['Molecules']})"
                      for _, r in pgroups.iterrows()))

    valid = validate(modes)
    if not valid.empty:
        bad = valid[valid["Verdict"] != "match"]
        print(f"  validation vs library_scores.csv: "
              f"{(valid['Verdict'] == 'match').sum()}/{len(valid)} match"
              + (f", {len(bad)} to review" if len(bad) else ""))

    out = Path(args.out)
    write(out, modes, refs, bonds, per_mol, pgroups, valid, failures)
    print(f"\nwrote {out}  ({out.stat().st_size/1e6:.2f} MB)")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
